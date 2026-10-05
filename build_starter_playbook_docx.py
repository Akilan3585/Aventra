from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"C:\Users\DELL\OneDrive\Desktop\Work\Projects\AventraAI\AventraAI\PXL-Starter-Playbook-Research-Report.docx"
NAVY = "15324B"
BLUE = "2563EB"
PALE = "F5F9FF"
GREY = "D9E2EC"


def shade(cell, colour):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), colour)
    tc_pr.append(shd)


def borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = tc_borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            tc_borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:color"), GREY)


def set_cell_margins(cell, top=100, start=115, bottom=100, end=115):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def no_table_borders(table):
    tbl_pr = table._tbl.tblPr
    tbl_borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement("w:" + edge)
        node.set(qn("w:val"), "nil")
        tbl_borders.append(node)
    tbl_pr.append(tbl_borders)


def style_run(run, size=12, bold=False, colour="000000"):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(colour)


def add_para(doc, text="", size=12, bold=False, colour="000000", before=0, after=8, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    if align:
        p.alignment = align
    style_run(p.add_run(text), size, bold, colour)
    return p


def add_title(doc, title, eyebrow=None, lead=None):
    if eyebrow:
        add_para(doc, eyebrow.upper(), 9, True, BLUE, after=10)
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(9)
    style_run(p.add_run(title), 29, True, "000000")
    if lead:
        add_para(doc, lead, 14, False, "3F5E79", after=18)


def add_heading(doc, title):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(10)
    style_run(p.add_run(title), 18, True, "000000")


def add_flow(doc, items):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(18)
    for i, item in enumerate(items):
        style_run(p.add_run(item), 12, True, NAVY)
        if i < len(items) - 1:
            style_run(p.add_run("  ->  "), 12, True, BLUE)


def add_comparison_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, header in enumerate(headers):
        cell = hdr.cells[i]
        cell.width = Inches(widths[i])
        shade(cell, NAVY)
        borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style_run(p.add_run(header), 8.5, True, "FFFFFF")
    for r_index, row_data in enumerate(rows):
        row = table.add_row()
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            cell.width = Inches(widths[i])
            borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_index % 2 == 1:
                shade(cell, PALE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.04
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in (2, 3) else WD_ALIGN_PARAGRAPH.LEFT
            style_run(p.add_run(value), 8.2, i == 0, NAVY if i == 0 else "26394A")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_two_columns(doc, items):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    no_table_borders(table)
    for pair_start in range(0, len(items), 2):
        row = table.add_row()
        for col in range(2):
            cell = row.cells[col]
            cell.width = Inches(4.0)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell, 80, 110, 120, 110)
            if pair_start + col >= len(items):
                continue
            heading, body = items[pair_start + col]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(5)
            style_run(p.add_run(heading), 12, True, NAVY)
            p2 = cell.add_paragraph()
            p2.paragraph_format.space_after = Pt(12)
            p2.paragraph_format.line_spacing = 1.1
            style_run(p2.add_run(body), 10.5, False, "3F5E79")
    return table


def next_page(doc):
    doc.add_page_break()


doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width = Inches(11)
section.page_height = Inches(8.5)
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.5)
section.left_margin = Inches(0.6)
section.right_margin = Inches(0.6)

normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)

# Slide 1
add_title(doc, "PXL Starter Playbook Research Report", "Research report | product direction", "A concise view of competitor learning models, learner needs and the product direction for Starter Playbook.")
add_para(doc, "By Akilan and Muhammed Umer", 13, True, NAVY, after=22)
add_heading(doc, "Research link")
add_para(doc, "https://app.notion.com/p/3e9a83de42f781b2a3f9f008edf2b157?pvs=204", 10, False, BLUE, after=20)
add_heading(doc, "Research direction")
add_para(doc, "Help learners move from understanding a tool to using it with confidence.", 17, True, NAVY, after=11)
add_para(doc, "The strongest signal from the research is clear direction + short structured learning + practical tasks + real-world application + proof that learners can do the work independently.", 13, False, "3F5E79")

# Slide 2
next_page(doc)
add_title(doc, "Competitor report", "Competitor landscape", "The first group shows how learning models, practice and paid value differ across platforms.")
add_comparison_table(doc,
    ["Platform", "Main focus / learning style", "Practical exposure", "Free access", "Paid value and takeaway"],
    [
        ["KodeKloud", "DevOps, cloud and AI paths, courses and labs.", "High", "Limited introductory content.", "Standard, Pro and AI plans. Labs are a strong benchmark; labs alone are not a differentiator."],
        ["Udemy", "Broad instructor marketplace, mainly video-led.", "Low-Medium", "Some free courses.", "One-time purchase or Personal Plan. Range is large, but course quality and practice vary."],
        ["DataCamp", "Data and AI courses with interactive browser lessons.", "High", "First chapter of every course.", "Premium: $28/month annually listed. Frequent exercises set a high expectation for practice."],
    ], [1.05, 1.85, 0.95, 1.2, 4.15])

# Slide 3
next_page(doc)
add_title(doc, "Competitor report", "Competitor landscape", "The second group is more focused on structured career learning, certificates and projects.")
add_comparison_table(doc,
    ["Platform", "Main focus / learning style", "Practical exposure", "Free access", "Paid value and takeaway"],
    [
        ["NxtWave", "Career learning with projects, mentoring and placement support.", "Medium-High", "Varies by program.", "Launchpad: INR 11,999/year listed. A strong role roadmap; prices differ by program."],
        ["Coursera", "University and company courses, certificates and degrees.", "Medium", "Free/audit access varies.", "$49-$79/month for many programs. Strong credibility; completion does not prove independent skill."],
        ["Udacity", "Career programs with projects, reviews and support.", "High", "Free course content.", "$249/month listed. Projects help, but feedback quality and freshness matter."],
    ], [1.05, 1.85, 0.95, 1.2, 4.15])
add_para(doc, "Shared pattern: Free introductory content lowers the barrier to try a platform. Paid plans promise deeper paths, more practice, projects, feedback or career support.", 11, True, NAVY, before=15, after=8)
add_para(doc, "Starter Playbook implication: A free entry point can work, provided the paid experience leads to a clearer practical outcome than 'more content.'", 11, False, "3F5E79")

# Slide 4
next_page(doc)
add_title(doc, "Make the learning tool-specific", "Our main direction", "Learners should see the right explanation, then use the tool in the same session. The content format can change, but the practice should stay tied to the tool being learned.")
add_flow(doc, ["Short video or slides", "Tool walkthrough", "Guided practice", "Tool-specific check"])
add_two_columns(doc, [
    ("1. Teach the task clearly", "Use short videos, slides, diagrams or a live walkthrough. The format should fit the topic and avoid turning every lesson into a long recording."),
    ("2. Keep the tool in view", "Show what the learner needs to do inside the relevant tool. A Git lesson should use Git workflows; a cloud lesson should use the cloud interface or lab."),
    ("3. Check understanding soon after", "Ask questions or give a small activity while the steps and decisions are still fresh."),
    ("4. Build toward real work", "Later modules can combine several tools into a realistic task. The first step is still a clear tool-specific foundation."),
])

# Slide 5
next_page(doc)
add_title(doc, "Give learners a clear finish line", "What we will provide", "Each tool-specific unit should show what the learner will receive, how progress is checked and when a certificate is issued.")
add_flow(doc, ["Learn with video or slides", "Practice in the tool", "Complete tool-specific quiz", "Finish required check", "Receive certificate"])
add_two_columns(doc, [
    ("What we provide", "Short explanations, guided practice, tool-specific quizzes and clear feedback. Selected paths can also include a practical check."),
    ("How certification works", "A certificate is issued after the learner completes the required content and checks for that unit. It records completion and assessed activities; it should not claim to guarantee job readiness."),
    ("For the learner", "They can see the route from the first lesson to a completed, documented unit."),
    ("For the product", "The completion rules stay consistent even when content is delivered through videos, slides or different tool environments."),
])

# Slide 6
next_page(doc)
add_title(doc, "How do we evaluate learners?", "Evaluation", "There are two possible approaches. Both can be useful, but they work at different stages of the product.")
add_two_columns(doc, [
    ("Approach 1: Make learners complete a task and check that", "Give a task in the relevant tool, then validate the work or review the output. This gives stronger evidence that the learner can do the work."),
    ("Approach 2: Interactive tool-specific quizzes", "Ask questions linked to the video, the tool and the decisions the learner needs to make. This keeps learners active and is easier to apply across many tools."),
    ("Task completion", "Best evidence of practical ability, but it needs tool environments, checking logic and support when learners get stuck."),
    ("Tool-specific quizzes", "Easy to standardize across tools, quick to complete and useful for checking understanding before moving to larger practical work."),
])

# Slide 7
next_page(doc)
add_title(doc, "Task completion gives stronger evidence, but it is costly to scale", "Approach 1", "Practical tasks should stay in the product, but they need to be used where they add the most value.")
add_two_columns(doc, [
    ("What it looks like", "The learner completes a task in a tool. The product checks the result, asks for an output or gives a reviewer enough evidence to confirm the work."),
    ("Why it becomes difficult", "Starter Playbook will be tool-specific and many tools can be added. Every tool may need its own environment, validation rules, task setup and help flow."),
    ("Use it where it matters most", "Keep practical checks for important milestones or end-of-unit work, where proof of ability is more valuable."),
    ("Do not make every lesson a lab", "A lab for every small concept will make the product slower and harder to maintain. Short quizzes can keep the learning active between practical checks."),
])

# Slide 8
next_page(doc)
add_title(doc, "Start with one tool-specific learning unit", "Proposed first version", "Test the smallest complete experience before adding a large library of tools or complex task validation.")
add_flow(doc, ["5-10 min explanation", "Interactive quiz", "Guided tool practice", "One practical check"])
add_two_columns(doc, [
    ("Test first", "Use one tool and one learner outcome. See whether learners can follow the content, stay engaged and complete the practical check."),
    ("Decide later", "If learners need more proof of ability, add stronger task validation only where the value is clear. Do not build every tool environment before the learning flow is validated."),
])
add_para(doc, "Research link: https://app.notion.com/p/3e9a83de42f781b2a3f9f008edf2b157?pvs=204", 10, False, BLUE, before=17, after=0)

doc.save(OUT)
print(OUT)
